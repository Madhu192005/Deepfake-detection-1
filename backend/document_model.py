import fitz
import docx
import pytesseract
from PIL import Image
import cv2
import numpy as np
import io
import re

pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'

AI_PATTERNS = [
    r'\bcertainly\b', r'\bof course\b', r'\bsure\b', r'\bas an ai\b',
    r'\bi cannot\b', r'\bi am unable\b', r'\bit is important to note\b',
    r'\bit is worth noting\b', r'\bin conclusion\b', r'\bfurthermore\b',
    r'\bmoreover\b', r'\bnevertheless\b', r'\bin summary\b', r'\bto summarize\b',
    r'\boverall\b', r'\bin addition\b', r'\bfirst and foremost\b',
    r'\bnotably\b', r'\bsignificantly\b', r'\bin this context\b',
    r'\bit should be noted\b', r'\bultimately\b',
]

def analyze_ai_text(text: str) -> dict:
    if not text or len(text.strip()) < 50:
        return {"ai_score": 0, "verdict": "INSUFFICIENT TEXT", "word_count": 0, "ai_patterns_found": [], "details": "Not enough text"}
    text_lower = text.lower()
    words = text.split()
    word_count = len(words)
    found_patterns = [p.replace(r'\b','').replace('\\','') for p in AI_PATTERNS if re.search(p, text_lower)]
    sentences = [s.strip() for s in re.split(r'[.!?]+', text) if len(s.strip()) > 10]
    avg_sentence_len = 0
    sentence_variance = 0
    if sentences:
        lengths = [len(s.split()) for s in sentences]
        avg_sentence_len = sum(lengths) / len(lengths)
        if len(lengths) > 1:
            mean = avg_sentence_len
            sentence_variance = (sum((x-mean)**2 for x in lengths)/len(lengths))**0.5
    ai_score = 0
    ai_score += min(len(found_patterns) * 8, 40)
    if sentences and sentence_variance < 5: ai_score += 30
    elif sentences and sentence_variance < 10: ai_score += 15
    if word_count > 200: ai_score += 10
    ai_score = min(ai_score, 100)
    if ai_score >= 60: verdict = "LIKELY AI-GENERATED"
    elif ai_score >= 35: verdict = "POSSIBLY AI-GENERATED"
    else: verdict = "LIKELY HUMAN-WRITTEN"
    return {
        "ai_score": round(ai_score, 2), "verdict": verdict,
        "word_count": word_count, "ai_patterns_found": found_patterns[:5],
        "avg_sentence_length": round(avg_sentence_len, 1),
        "details": f"Found {len(found_patterns)} AI patterns in {word_count} words"
    }

def analyze_metadata(metadata: dict) -> dict:
    issues = []
    tamper_score = 0
    creator = metadata.get("creator","") or ""
    producer = metadata.get("producer","") or ""
    author = metadata.get("author","") or ""
    created = metadata.get("creationDate","") or metadata.get("created","") or ""
    modified = metadata.get("modDate","") or metadata.get("modified","") or ""
    if created and modified and created != modified:
        issues.append("Document was modified after creation")
        tamper_score += 20
    for tool in ["photoshop","gimp","inkscape","paint","editor"]:
        if tool in producer.lower() or tool in creator.lower():
            issues.append(f"Edited with: {tool}")
            tamper_score += 25
    if not author and not creator:
        issues.append("Missing author/creator metadata")
        tamper_score += 15
    return {
        "tamper_score": min(tamper_score, 100), "issues": issues,
        "metadata": {"author": author or "Unknown", "creator": creator or "Unknown",
                     "producer": producer or "Unknown",
                     "created": str(created)[:20] if created else "Unknown",
                     "modified": str(modified)[:20] if modified else "Unknown"}
    }

def analyze_pdf(file_bytes: bytes) -> dict:
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        full_text = "".join(page.get_text() for page in doc)
        metadata = doc.metadata or {}
        page_count = doc.page_count
        fonts = set(font[3] for page in doc for font in page.get_fonts())
        font_count = len(fonts)
        doc.close()
        ai_result = analyze_ai_text(full_text)
        meta_result = analyze_metadata(metadata)
        tamper_score = meta_result["tamper_score"]
        if font_count > 5:
            tamper_score = min(tamper_score + 20, 100)
            meta_result["issues"].append(f"Unusual number of fonts: {font_count}")
        overall_fake_score = (ai_result["ai_score"] * 0.5) + (tamper_score * 0.5)
        return {
            "file_type": "PDF", "page_count": page_count, "text_extracted": len(full_text) > 0,
            "ai_analysis": ai_result,
            "tamper_analysis": {"tamper_score": round(tamper_score,2), "issues": meta_result["issues"],
                                "font_count": font_count, "metadata": meta_result["metadata"]},
            "overall_fake_score": round(overall_fake_score, 2),
            "overall_verdict": "FAKE/SUSPICIOUS" if overall_fake_score >= 50 else "LIKELY AUTHENTIC"
        }
    except Exception as e:
        return {"error": f"PDF analysis failed: {str(e)}"}

def analyze_docx(file_bytes: bytes) -> dict:
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = "\n".join(para.text for para in doc.paragraphs)
        props = doc.core_properties
        metadata = {
            "author": str(props.author or ""), "created": str(props.created or ""),
            "modified": str(props.modified or ""), "last_modified_by": str(props.last_modified_by or ""),
            "revision": str(props.revision or "")
        }
        fonts = set(run.font.name for para in doc.paragraphs for run in para.runs if run.font.name)
        font_count = len(fonts)
        try: revision = int(props.revision or 0)
        except: revision = 0
        ai_result = analyze_ai_text(full_text)
        tamper_issues = []
        tamper_score = 0
        if props.last_modified_by and props.author and str(props.last_modified_by) != str(props.author):
            tamper_issues.append(f"Modified by different person: {props.last_modified_by}")
            tamper_score += 25
        if revision > 20:
            tamper_issues.append(f"High revision count: {revision}")
            tamper_score += 20
        if font_count > 6:
            tamper_issues.append(f"Too many fonts: {font_count}")
            tamper_score += 15
        overall_fake_score = (ai_result["ai_score"] * 0.5) + (tamper_score * 0.5)
        return {
            "file_type": "DOCX", "paragraph_count": len(doc.paragraphs), "text_extracted": len(full_text) > 0,
            "ai_analysis": ai_result,
            "tamper_analysis": {"tamper_score": round(tamper_score,2), "issues": tamper_issues,
                                "font_count": font_count, "revision_count": revision, "metadata": metadata},
            "overall_fake_score": round(overall_fake_score, 2),
            "overall_verdict": "FAKE/SUSPICIOUS" if overall_fake_score >= 50 else "LIKELY AUTHENTIC"
        }
    except Exception as e:
        return {"error": f"DOCX analysis failed: {str(e)}"}

def analyze_image_document(image: Image.Image) -> dict:
    try:
        extracted_text = pytesseract.image_to_string(image)
        img_array = np.array(image)
        img_bgr = cv2.cvtColor(img_array, cv2.COLOR_RGB2BGR)
        img_gray = cv2.cvtColor(img_bgr, cv2.COLOR_BGR2GRAY)
        buffer = io.BytesIO()
        image.save(buffer, format='JPEG', quality=75)
        buffer.seek(0)
        compressed = Image.open(buffer)
        original_array = np.array(image.convert('RGB'))
        compressed_array = np.array(compressed.convert('RGB'))
        if original_array.shape != compressed_array.shape:
            compressed_array = cv2.resize(compressed_array, (original_array.shape[1], original_array.shape[0]))
        ela_diff = cv2.absdiff(original_array, compressed_array)
        ela_score = float(np.mean(ela_diff))
        laplacian_var = cv2.Laplacian(img_gray, cv2.CV_64F).var()
        tamper_issues = []
        tamper_score = 0
        if ela_score > 15:
            tamper_issues.append(f"High ELA score: possible editing ({ela_score:.1f})")
            tamper_score += min(ela_score * 2, 50)
        if laplacian_var < 50:
            tamper_issues.append("Low sharpness: possible fake scan")
            tamper_score += 20
        tamper_score = min(tamper_score, 100)
        ai_result = analyze_ai_text(extracted_text)
        overall_fake_score = (ai_result["ai_score"] * 0.4) + (tamper_score * 0.6)
        return {
            "file_type": "IMAGE", "text_extracted": len(extracted_text.strip()) > 0,
            "ocr_text_preview": extracted_text[:200].strip() if extracted_text else "No text found",
            "ai_analysis": ai_result,
            "tamper_analysis": {"tamper_score": round(tamper_score,2), "issues": tamper_issues,
                                "ela_score": round(ela_score,2), "sharpness": round(laplacian_var,2)},
            "overall_fake_score": round(overall_fake_score, 2),
            "overall_verdict": "FAKE/SUSPICIOUS" if overall_fake_score >= 50 else "LIKELY AUTHENTIC"
        }
    except Exception as e:
        return {"error": f"Image analysis failed: {str(e)}"}